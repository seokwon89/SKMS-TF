"""SKI 템플릿 양식(A4 가로, 나눔스퀘어, 남색 띠)으로 SK이노베이션 계열 Cold Eye Review 덱을 만든다.
템플릿의 마스터·테마·슬라이드 크기만 쓰고 기존 슬라이드는 모두 지운다."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

# 기존 슬라이드 삭제
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

LAYOUT = prs.slide_layouts[1]
W = prs.slide_width
# 템플릿 실측 색상 (테마 dk2=0E2841, lt1=FFFFFF 기준 lumMod/lumOff 계산값 포함)
NAVY = RGBColor(0x00, 0x20, 0x60)    # 제목 띠: srgbClr 002060
ACCENT = RGBColor(0x21, 0x5F, 0x9A)  # 채움 태그·번호 배지·강조 글자: tx2 lumMod75 lumOff25
INK = RGBColor(0x00, 0x00, 0x00)     # 본문: tx1 / prstClr black
SUB = RGBColor(0x59, 0x59, 0x59)     # 보조 본문: tx1 lumMod65 lumOff35
GRAY = RGBColor(0x80, 0x80, 0x80)    # 머리말 섹션명·각주: bg1 lumMod50
LGRAY = RGBColor(0xA6, 0xA6, 0xA6)   # 열 제목 밑줄·간지 라벨: bg1 lumMod65
PILLLINE = RGBColor(0xBF, 0xBF, 0xBF)  # 테두리 태그 선: bg1 lumMod75
PILLTXT = RGBColor(0x40, 0x40, 0x40)   # 테두리 태그 글자: tx1 lumMod75 lumOff25
FILL = RGBColor(0xF2, 0xF2, 0xF2)    # 표 머리·숫자 상자: bg1 lumMod95
MIDGRAY = RGBColor(0xD9, 0xD9, 0xD9) # 중립 상태 칸: bg1 lumMod85
LBLUE = RGBColor(0xDC, 0xEA, 0xF7)   # 연한 강조 상자: tx2 lumMod10 lumOff90
BAR = RGBColor(0xA1, 0xC4, 0xE3)     # 차트 막대: accent2(629DD1, themeOverride) lumMod60 lumOff40
RED = RGBColor(0xC0, 0x00, 0x00)     # 부정 수치·경고: srgbClr C00000
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = LGRAY
BLUE = BAR
ORANGE = RED
F_R, F_B, F_X = "나눔스퀘어", "나눔스퀘어 Bold", "나눔스퀘어 ExtraBold"


def set_font(run, face, size, color=INK, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)
    rPr.set("spc", "-30")


def tb(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.04, name=None):
    """paras: list of paragraphs; paragraph = list of (text, face, size, color, bold) or dict with 'runs', 'bullet', 'level', 'space'"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    first = True
    for p in paras:
        if isinstance(p, list):
            p = {"runs": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", align)
        if p.get("space") is not None:
            para.space_after = Pt(p["space"])
        if p.get("before") is not None:
            para.space_before = Pt(p["before"])
        pPr = para._p.get_or_add_pPr()
        if p.get("bullet"):
            ind = p.get("indent", 0.16)
            pPr.set("marL", str(int(Inches(ind + 0.16 * p.get("level", 0)))))
            pPr.set("indent", str(-int(Inches(0.16))))
            bu = etree.SubElement(pPr, qn("a:buFont")); bu.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", p["bullet"])
        else:
            etree.SubElement(pPr, qn("a:buNone"))
        for (t, face, size, color, bold) in p["runs"]:
            r = para.add_run()
            r.text = t
            set_font(r, face, size, color, bold)
    return box


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75, name=None, radius=0.16667):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius  # 템플릿: 채움 태그·배지 기본값 0.16667, 테두리 태그 0.19092
    s.text_frame.text = ""
    return s


def hline(slide, x, y, w, color=LINE, weight=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(weight)
    return c


def new_slide():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


page = [0]


SLDNUM = next(ph for ph in LAYOUT.placeholders if ph.placeholder_format.type == 13)  # PP_PLACEHOLDER.SLIDE_NUMBER


def pageno(s):
    """레이아웃의 슬라이드 번호 개체틀을 그대로 복제해 템플릿 서식(위치·크기·색)을 쓴다."""
    page[0] += 1
    el = copy.deepcopy(SLDNUM._element)
    s.shapes._spTree.append(el)


def header(s, section, sub):
    tb(s, 0.3, 0.2, 9.8, 0.42, [[(section + "  ", F_X, 18, GRAY, False), ("|  ", F_X, 18, INK, False), (sub, F_X, 18, INK, False)]], name="Header", margin=0)


def headline(s, text, y=0.82):
    tb(s, 0.3, y, 10.23, 0.5, [[(text, F_X, 14, INK, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Headline")


def source(s, text):
    tb(s, 0.3, 6.98, 7.2, 0.4, [[(text, F_R, 7.5, GRAY, False)]], name="Source", anchor=MSO_ANCHOR.BOTTOM)


def pill(s, x, y, w, text, filled=False):
    if filled:
        r = rect(s, x, y, w, 0.24, fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        col = WHITE
    else:
        r = rect(s, x, y, w, 0.24, fill=WHITE, line=PILLLINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.19092)
        col = PILLTXT
    tb(s, x, y, w, 0.24, [[(text, F_B, 9.5, col, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)


def col_title(s, x, w, y, text):
    tb(s, x, y, w, 0.32, [[(text, F_X, 13, INK, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, x, y + 0.36, w, LINE, 1.0)


STATUS = {"ok": ("진전", ACCENT, WHITE), "mid": ("미확인", MIDGRAY, INK), "no": ("미달·괴리", RED, WHITE)}

# ---------------------------------------------------------------- slides
# 1. Cover
s = new_slide()
rect(s, 0.6, 2.05, 9.63, 0.95, fill=NAVY, name="Title band")
tb(s, 0.6, 2.05, 9.63, 0.95, [[("SK이노베이션 계열  Cold Eye Review", F_X, 30, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Title")
toc = [
    {"runs": [("총평 및 부문별 성적표", F_R, 12, INK, False)], "bullet": "■", "space": 6},
    {"runs": [("사업부문별 리뷰: 회사 목표 → 목표 대비 판정 → CEO 최우선 과제", F_R, 12, INK, False)], "bullet": "■", "space": 3},
    {"runs": [("석유  ·  화학  ·  윤활유  ·  배터리  ·  E&S", F_R, 11, INK, False)], "bullet": "□", "level": 1, "space": 6},
    {"runs": [("그룹 재무, 다음 분기 체크포인트", F_R, 12, INK, False)], "bullet": "■", "space": 6},
]
tb(s, 0.8, 3.2, 8.5, 1.6, toc, name="Contents")
tb(s, 0.6, 6.25, 9.63, 0.3, [[("’26. 10", F_R, 11, INK, False)]], align=PP_ALIGN.CENTER)
tb(s, 0.6, 6.55, 9.63, 0.3, [[("기준: 2025 사업보고서(2025.12.31) 및 2026년 공개 자료 · 투자 권유 아님", F_R, 8.5, GRAY, False)]], align=PP_ALIGN.CENTER)
pageno(s)

# 2. Executive summary
s = new_slide()
header(s, "총평", "부문별 성적표")
headline(s, "정유는 펌프, 윤활유·E&S는 금고, 배터리는 블랙홀. 2026년 반등은 유가·보상금 효과를 걷어내고 봐야 한다")
kp = [("80.3조", "연결 매출 ('25)", INK), ("0.45조", "연결 영업이익 (이익률 0.6%)", INK), ("−3.35조", "지배주주 순손실 ('25)", RED), ("5.6조", "'26 상반기 영업이익 (보도)", ACCENT)]
for i, (v, l, c) in enumerate(kp):
    x = 0.3 + i * 2.58
    rect(s, x, 1.45, 2.45, 0.85, fill=FILL)
    tb(s, x, 1.5, 2.45, 0.45, [[(v, F_X, 20, c, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x, 1.95, 2.45, 0.3, [[(l, F_R, 9, GRAY, False)]], align=PP_ALIGN.CENTER)
rows = [
    ("부문", "주 법인", "'25 매출", "'25 영업이익", "이익률", "목표 판정", "한 줄 평"),
    ("석유", "SK에너지, SK온(트레이딩)", "47.2", "−0.17*", "−0.4%*", "진전1 · 미확인2 · 미달1", "매출은 착시, 이익은 유가"),
    ("화학", "SK지오센트릭", "8.9", "−0.15*", "−1.3%*", "진전1 · 미확인2 · 미달2", "덜 잃는 게 목표"),
    ("윤활유", "옛 SK엔무브(현 SK온)", "3.8", "미공시", "—", "진전1 · 미확인2", "최고의 사업, 이상한 위치"),
    ("배터리", "SK온", "7.0", "−0.93", "−13.3%", "진전2 · 미확인1 · 미달1", "흑자의 품질을 따져야"),
    ("E&S", "SK이노베이션 E&S CIC", "11.9", "0.68", "5.7%", "진전1 · 미확인2", "방파제"),
    ("그룹", "SK이노베이션 연결", "80.3", "0.45", "0.6%", "진전2 · 미확인1 · 미달1", "덩치는 컸고 남은 건 적었다"),
]
cw = [0.85, 2.05, 0.8, 1.05, 0.8, 1.95, 2.73]
tbl = s.shapes.add_table(len(rows), len(cw), Inches(0.3), Inches(2.55), Inches(sum(cw)), Inches(0.36 * len(rows))).table
for j, w_ in enumerate(cw):
    tbl.columns[j].width = Inches(w_)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.4)
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = FILL if i == 0 else WHITE
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT if j in (2, 3, 4) and i else (PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT)
        r = p.add_run(); r.text = val
        col = INK
        if i and j in (3, 4) and val.startswith("−"):
            col = RED
        set_font(r, F_B if (i == 0 or j == 0) else F_R, 10 if i else 10, col, i == 0 or j == 0)
tbl_el = s.shapes[-1]._element
tb(s, 0.3, 5.5, 10.2, 0.9, [
    {"runs": [("매출: 석유·화학·윤활유는 SK이노베이션 연결 기준(내부거래 제거 후), 배터리는 SK온 사업부문, E&S는 SK이노베이션 사업부문 기준. 단위 조원.", F_R, 8.5, GRAY, False)], "space": 2},
    {"runs": [("* 석유는 부문 영업이익이 따로 공시되지 않아 SK에너지 연결 수치, 화학은 SK지오센트릭 연결 수치(매출 11.5조 기준)를 표시.", F_R, 8.5, GRAY, False)], "space": 2},
    {"runs": [("목표 판정: 회사가 사업보고서·주주총회·컨퍼런스콜에서 밝힌 목표를 실제 숫자와 대조한 결과(세부는 부문별 페이지).", F_R, 8.5, GRAY, False)]},
])
source(s, "출처: SK이노베이션·SK에너지·SK온·SK지오센트릭 2025 사업보고서, 2026년 2분기 실적 보도(Investing.com, 더트래커)")
pageno(s)

# 3. Divider
def divider(top, title):
    s = new_slide()
    tb(s, 1.6, 2.55, 7.6, 0.4, [[(top, F_X, 16, LGRAY, True)]])
    rect(s, 1.6, 2.98, 7.63, 0.78, fill=NAVY)
    tb(s, 1.6, 2.98, 7.63, 0.78, [[(title, F_X, 28, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="Title")
    pageno(s)

divider("SK이노베이션 계열  Cold Eye Review", "사업부문별 리뷰")

# Segment slide template
def segment(name, headline_txt, stats, goals, reviews, ceos, msg, pull, src, section="사업부문별 리뷰"):
    s = new_slide()
    header(s, section, name)
    # number badge before headline? template uses badge in header; keep header text
    headline(s, headline_txt)
    LX, LW, RX, RW = 0.3, 4.95, 5.55, 4.98
    col_title(s, LX, LW, 1.4, "검토 내용")
    col_title(s, RX, RW, 1.4, "판정 및 CEO 과제")
    # stats
    sw = (LW - 0.2) / 3
    for i, (v, l, neg) in enumerate(stats):
        x = LX + i * (sw + 0.1)
        rect(s, x, 1.95, sw, 0.82, fill=FILL)
        tb(s, x, 1.98, sw, 0.42, [[(v, F_X, 17, RED if neg else ACCENT, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + 0.04, 2.38, sw - 0.08, 0.38, [[(l, F_R, 8.5, SUB, False)]], align=PP_ALIGN.CENTER)
    pill(s, LX, 2.95, 1.35, "회사가 내건 목표")
    gp = []
    for t, src_ in goals:
        gp.append({"runs": [(t, F_R, 11, INK, False)], "bullet": "•", "space": 1})
        gp.append({"runs": [(src_, F_R, 8, GRAY, False)], "space": 9 if len(goals) <= 3 else 6, "indent": 0.16, "bullet": None})
    for p in gp:
        if p.get("bullet") is None:
            p.pop("bullet")
    box = tb(s, LX, 3.27, LW, 2.55, gp)
    # indent source lines
    for idx, para in enumerate(box.text_frame.paragraphs):
        if idx % 2 == 1:
            para._p.get_or_add_pPr().set("marL", str(int(Inches(0.16))))
    # pull quote
    rect(s, LX, 6.08, LW, 0.85, fill=LBLUE)
    tb(s, LX + 0.15, 6.08, LW - 0.3, 0.85, [[("“ " + pull + " ”", F_B, 11.5, ACCENT, True)]], anchor=MSO_ANCHOR.MIDDLE)
    # right: review table
    pill(s, RX, 1.95, 1.35, "목표 대비 판정", filled=True)
    n = len(reviews)
    rh = min(0.55, 2.1 / n)
    t = s.shapes.add_table(n, 3, Inches(RX), Inches(2.28), Inches(RW), Inches(rh * n)).table
    for j, w_ in enumerate([1.3, 0.78, 2.9]):
        t.columns[j].width = Inches(w_)
    for i, (goal, st, why) in enumerate(reviews):
        t.rows[i].height = Inches(rh)
        lab, bg, fg = STATUS[st]
        vals = [(goal, F_B, 10, INK, True, WHITE), (lab, F_B, 9, fg, True, bg), (why, F_R, 9.5, INK, False, WHITE)]
        for j, (v, face, sz, col, b, fillc) in enumerate(vals):
            c = t.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = fillc
            c.margin_left = c.margin_right = Inches(0.05); c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
            p.alignment = PP_ALIGN.CENTER if j == 1 else PP_ALIGN.LEFT
            r = p.add_run(); r.text = v
            set_font(r, face, sz, col, b)
    ty = 2.28 + rh * n + 0.15
    pill(s, RX, ty, 1.35, "내가 CEO라면", filled=True)
    cp = [{"runs": [(f"{k+1}. ", F_X, 10.5, ACCENT, True), (c_, F_R, 10.5, INK, False)], "space": 5} for k, c_ in enumerate(ceos)]
    tb(s, RX, ty + 0.32, RW, 6.0 - (ty + 0.32), cp)
    rect(s, RX, 6.08, RW, 0.85, fill=ACCENT)
    tb(s, RX + 0.15, 6.08, RW - 0.3, 0.85, [{"runs": [("CEO 메시지", F_B, 9, LBLUE, True)], "space": 3}, [(msg, F_X, 11.5, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    source(s, src)
    pageno(s)
    return s

SEGS = [
 ("01  석유", "47.2조를 팔지만 이익은 유가가 정한다. 정유에서 볼 숫자는 매출이 아니라 마진이다",
  [("47.2조", "계열 석유사업 외부 매출('25)", False), ("−0.17조", "SK에너지 영업이익('25)", True), ("0.41→−0.17", "SK에너지 영업이익 '23→'25", True)],
  [("‘동북아 No.1 생산/마케팅 경쟁력’의 Energy Provider로 진화", "SK에너지 2025 사업보고서"),
   ("Operation Improvement로 수익구조 혁신, 생존력 강화", "SK에너지 2025 사업보고서"),
   ("SAF 연산 10만톤 생산체계, '25.1 국내 정유사 최초 유럽 수출", "SK에너지 2025 사업보고서"),
   ("트레이딩: 바이오·SAF 거래 확대, 배터리 원소재 트레이딩 진출 예정", "SK이노베이션 2025 사업보고서")],
  [("동북아 No.1", "no", "영업이익 0.41→0.05→−0.17조, 3년 연속 하락"),
   ("Operation Improvement", "mid", "절감 성과를 수치로 공시하지 않음"),
   ("SAF·저탄소 제품", "ok", "유럽 수출·인천공항 공급. 이익 기여는 미공시"),
   ("원소재 트레이딩", "mid", "‘진출 예정’ 단계, 실적 없음")],
  ["재고효과를 뺀 정제마진 이익을 분기마다 따로 공시", "Operation Improvement에 배럴당 원가 절감 목표·실적을 붙여 공개", "SAF 다년 오프테이크 계약으로 프리미엄 고정"],
  "유가가 준 이익은 유가가 가져간다. 우리가 통제하는 마진만 우리 실적이다",
  "재고이익은 실력이 아니라 유가가 준 선물이다. 선물은 돌려달라고 할 때가 온다",
  "출처: SK에너지·SK이노베이션 2025 사업보고서, 2026년 2분기 실적 보도(Investing.com, 에너지데일리)"),
]

segment(*SEGS[0])

# 5. Chart slide: 계열 석유·화학 매출 구조
s = new_slide()
header(s, "사업부문별 리뷰", "01  석유  ·  계열 매출 구조")
headline(s, "법인 매출을 더하면 134.6조, 내부거래 73.3조를 지우면 계열 밖으로 판 것은 61.3조. 그중 석유는 47.2조")
col_title(s, 0.3, 4.95, 1.4, "법인별 매출 (내부거래 포함, '25)")
col_title(s, 5.55, 4.98, 1.4, "SK이노베이션 연결 기준 (내부거래 제거 후)")
cd = CategoryChartData()
cd.categories = ["SK온 석유사업", "SK에너지", "SK지오센트릭(화학)", "SK온 윤활유(11~12월)", "그 밖의 계열·조정(잔여)"]
cd.add_series("매출(조원)", (49.0, 39.3, 11.5, 0.8, 34.1))
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.3), Inches(1.9), Inches(4.95), Inches(3.6), cd).chart
cd2 = CategoryChartData()
cd2.categories = ["석유사업", "화학사업", "윤활유사업", "기타(잔여)", "내부거래 제거"]
cd2.add_series("매출(조원)", (47.2, 8.9, 3.8, 1.4, 73.3))
ch2 = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5.55), Inches(1.9), Inches(4.98), Inches(3.6), cd2).chart
for c_, colors in ((ch, [ACCENT, ACCENT, BAR, BAR, PILLLINE]), (ch2, [ACCENT, BAR, BAR, PILLLINE, MIDGRAY])):
    c_.has_legend = False
    c_.has_title = False
    plot = c_.plots[0]
    plot.gap_width = 60
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0"조"'; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10); dl.font.name = F_B; dl.font.color.rgb = INK
    ser = plot.series[0]
    for k, col in enumerate(colors):
        pt = ser.points[k]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = col
    ca = c_.category_axis
    ca.reverse_order = True
    ca.tick_labels.font.size = Pt(9.5); ca.tick_labels.font.name = F_R; ca.tick_labels.font.color.rgb = INK
    ca.format.line.color.rgb = LINE
    va = c_.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = 85
    va.minimum_scale = 0
pts = [
    {"runs": [("같은 기름이 두 번 매출로 잡힌다. ", F_B, 10, INK, True), ("SK에너지가 정제한 제품을 SK온의 트레이딩 조직(옛 SK트레이딩인터내셔널)이 사서 다시 판다.", F_R, 10, INK, False)], "bullet": "•", "space": 4},
    {"runs": [("윤활유는 1년치가 연결에 있다. ", F_B, 10, INK, True), ("1~10월은 SK엔무브(SK이노베이션 자회사), 11~12월은 SK온 기유·윤활유 부문으로 잡혔다.", F_R, 10, INK, False)], "bullet": "•", "space": 4},
    {"runs": [("부문 영업이익은 에너지·화학 합산 1.12조만 공시된다. ", F_B, 10, INK, True), ("석유만의 이익은 따로 나오지 않는다.", F_R, 10, INK, False)], "bullet": "•"},
]
tb(s, 0.3, 5.7, 10.23, 1.2, pts)
source(s, "출처: SK이노베이션 2025 사업보고서(사업의 내용, 주석 36 영업부문 정보: 에너지 및 화학 총매출 134.6조, 조정 −73.3조), SK온·SK에너지·SK지오센트릭 2025 사업보고서. 잔여·기타는 차감 계산")
pageno(s)

SEGS2 = [
 ("02  화학", "팔수록 줄고, 정리할수록 아프다. 범용 석유화학 공급과잉 한가운데서 지금 목표는 ‘덜 잃는 것’",
  [("11.5조", "SK지오센트릭 매출('25, 3년 연속↓)", False), ("−0.15조", "영업이익('25)", True), ("−0.52조", "중단영업손실(손상 0.43조 포함)", True)],
  [("최적 생산운영·원가 절감으로 원가 경쟁력 강화", "SK지오센트릭 경영진단(SK이노베이션 2025 사업보고서)"),
   ("제품 고부가화와 마케팅 확장으로 구조적 혁신", "같은 자료"),
   ("폐플라스틱·폐타이어 리사이클 신사업으로 성장 기반 구축", "같은 자료"),
   ("비핵심 자산 매각: 해외 기능성소재 법인('25 결정), 닝보SK합성고무 80%('26.1)", "SK이노베이션 사업보고서 주석 41, 마이데일리")],
  [("원가 경쟁력", "no", "기초유화 영업손실 1,306억원('25)"),
   ("고부가화", "mid", "신규 설비 '25 가동, 비중 목표·실적 미공시"),
   ("리사이클 신사업", "mid", "방향은 반복 제시, 매출 규모 미공시"),
   ("비핵심 자산 정리", "ok", "매각 실행. 대가로 손상 4,266억원"),
   ("울산 NCC 재편", "no", "이해관계 차이로 지연(보도)")],
  ["울산 NCC 재편을 연내 결론(감축 규모·통합 상대 확정)", "고부가 제품 매출 비중을 숫자 목표로 공개", "리사이클 투자는 상업 가동·손익분기 확인 후 다음 단계 집행"],
  "버티는 회사가 아니라, 줄여서 사는 회사가 돼야 한다",
  "구조조정의 첫 청구서는 이미 왔다. 문제는 다음 청구서가 몇 장 더 남았느냐다",
  "출처: SK지오센트릭·SK이노베이션 2025 사업보고서(주석 41 중단영업), 마이데일리(2026.10.22)"),
 ("03  윤활유", "계열에서 가장 좋은 사업이 배터리 회사 안으로 이사 갔다. SK온 숫자는 배터리와 윤활유를 나눠 봐야 한다",
  [("3.84조", "윤활유 매출('25, SK이노베이션 연결)", False), ("0.69조", "'26 2Q 영업이익(보도)", False), ("'25.11", "SK온에 흡수합병", False)],
  [("데이터센터 액침냉각 기업 GRC 지분투자('22 2,500만$, '25 393만$ 추가)", "SK이노베이션 2025 사업보고서"),
   ("윤활기유로 EV 배터리·ESS·데이터센터 액침냉각 공동 개발·상업화", "SK이노베이션 2025 사업보고서")],
  [("본업 수익성", "ok", "'26 2Q 0.69조. 중동 경쟁사 공급 차질이라는 외부 요인 큼"),
   ("액침냉각 상업화", "mid", "공동 개발 단계, 상업 수주·매출 미공시"),
   ("SK온 편입 후 독립성", "mid", "윤활유 현금 사용 원칙 미공개")],
  ["‘지정학 프리미엄’과 구조적 이익을 나눠 투자자에게 설명", "액침냉각 첫 상업 계약 확보(‘협업’을 ‘매출’로)", "SK온 내 부문별 자본배분 원칙 공개(배터리 지원 한도 설정)"],
  "좋은 사업의 현금은 공짜가 아니다. 쓸 곳과 한도를 먼저 정한다",
  "좋은 사업을 나쁜 사업 옆에 두면 평균은 좋아 보인다. 평균은 아무것도 설명하지 않는다",
  "출처: SK이노베이션·SK온 2025 사업보고서, 2026년 2분기 실적 보도(Investing.com), 디일렉·블로터(합병 배경)"),
 ("04  배터리", "적자보다 무서운 건 손상차손. 영업손실 0.93조인데 순손실은 5.36조, 공장값을 장부에서 지웠다",
  [("−0.93조", "배터리 영업이익('25)", True), ("4.78조", "유형자산 손상차손('25)", True), ("242%", "SK온 부채비율('25말)", True)],
  [("ESS로 포트폴리오 다각화, 미국·신재생 연계 시장 중심 확장", "SK이노베이션 2025 사업보고서"),
   ("Cell 라인 효율화, ESS 전용 라인 확보로 ESS 매출 비중 확대", "SK이노베이션 2025 사업보고서"),
   ("'26 북미 중심 ESS 수주 20GWh 이상", "4Q25 실적 컨퍼런스콜(이투데이·데일리안)"),
   ("전기차 배터리는 수익성 위주로 선별 수주·생산", "'26.3 정기주주총회(글로벌이코노믹)")],
  [("ESS 수주 20GWh", "mid", "공개 계약 약 10GWh(네오볼타 9, 플랫아이언 1). 회사는 달성 ‘긍정적’('26.10)"),
   ("ESS 전용 라인", "ok", "포드 합작 종료로 확보한 테네시 공장 ESS 전환 방향"),
   ("수익성 위주 운영", "no", "'26 2Q 흑자 0.82조에 고객 보상금·IRA 세액공제 포함"),
   ("설비 효율화", "ok", "CAPEX 9.8→3.6조('23→'25), 대가로 손상 4.78조")],
  ["보상금·AMPC를 뺀 영업이익 흑자 시점을 공개 목표로 설정", "테네시 ESS 전환 일정·목표 가동률을 분기 보고", "헝가리 등 유휴 해외 설비 연내 결정(전환·매각·축소)"],
  "보조금 없이 남는 1원이 진짜 흑자다",
  "흑자 전환은 뉴스지만, 보상금과 세액공제를 뺀 흑자만이 실적이다",
  "출처: SK온 2025 사업보고서(유형자산 주석, 사업부문별 재무정보), 머니투데이·이투데이(합작 종결), 한국금융신문(ESS 수주)"),
 ("05  E&S", "합병의 진짜 이유는 손익계산서에 있다. E&S 이익이 배터리·소재 적자의 58%를 메웠다",
  [("11.86조", "E&S 매출('25)", False), ("0.68조", "영업이익(이익률 5.7%)", False), ("58%", "배터리·소재 적자 보전율", False)],
  [("Electrification 시대의 Total Energy Company, 발전자산 확보로 전기사업자 전환", "SK이노베이션 2025 사업보고서"),
   ("전력 다소비 산업 통합 서비스, AI 데이터센터 등 대형 수요처 전력 솔루션", "사업보고서, '26.3 주주총회"),
   ("Global LNG 밸류체인 확장", "SK이노베이션 2025 사업보고서"),
   ("2026년을 ‘전기사업자 포지셔닝 원년’으로", "'26.3 주주총회(헤럴드경제)")],
  [("안정적 이익 창출", "ok", "영업이익 0.68조, 그룹의 방파제 역할"),
   ("전기사업자 전환", "mid", "발전 용량(GW)·고객 수 등 측정 지표 미공시"),
   ("LNG 밸류체인 확장", "mid", "바로사 하반기 가동 기대. 보령LNG터미널 등 자산 정리 병행(보도)")],
  ["전기사업자 KPI(발전 용량·판매량·산업 고객 수) 분기 공시", "바로사 상업 가동과 장기 판매계약 확정", "AI 데이터센터 전력 솔루션 첫 대형 계약 확보"],
  "방파제를 성장엔진이라 부르려면, 숫자로 된 이정표가 필요하다",
  "E&S는 성장 스토리가 아니라 방파제다. 파도가 계속 높으면 결국 넘친다",
  "출처: SK이노베이션 2025 사업보고서(사업부문별 재무정보), 헤럴드경제(주총), 데일리안(컨퍼런스콜), 뉴스웨이(자산 매각)"),
]
for seg in SEGS2:
    segment(*seg)

divider("SK이노베이션 계열  Cold Eye Review", "그룹 재무 및 시사점")

segment("06  그룹 재무", "자본의 40%는 남의 몫. 리밸런싱은 진행 중이지만 부채비율은 3년째 올랐다",
  [("190%", "연결 부채비율('25말)", True), ("22.5조", "순차입금(회사 발표, '25말)", False), ("−3.08조", "FCF('25)", True)],
  [("포트폴리오 리밸런싱을 빠른 시일 내 완수, 운영개선 고도화", "'26.1.2 신년사(머니투데이)"),
   ("재무구조 안정화: 핵심 사업 집중, 순차입금 감축", "'26.3.24 정기주주총회(글로벌이코노믹)"),
   ("2026년을 ‘체질 정상화’의 해로, 무배당·비핵심 자산 매각 지속", "공개 컨퍼런스콜 보도(파이낸스스코프, 딜사이트)")],
  [("순차입금 감축", "ok", "28.8→22.5조('25.9말→12말, 회사 발표 보도)"),
   ("신용도", "ok", "무디스 Ba1 전망 ‘안정적’ 상향('26). '27~'28 레버리지 재상승 전망"),
   ("리밸런싱 완수", "mid", "매각은 진행, 울산 석유화학 재편 지연"),
   ("부채비율", "no", "169→179→190%('23→'25)")],
  ["순차입금 목표를 ‘○년까지 ○조’로 숫자·기한 공표", "부문별 투하자본이익률(ROIC) 공시", "비지배지분(자본의 40%) 구조 단순화 로드맵 제시"],
  "리밸런싱의 끝은 매각 공시가 아니라, 주주 몫 이익이 늘어난 손익계산서다",
  "2026년 이익 반등이 재무를 고치는 건 맞다. 다만 그 반등의 상당 부분은 중동 리스크가 깔아 준 것이다",
  "출처: SK이노베이션 2025 사업보고서(요약연결재무정보·현금흐름표), 이투데이(무디스), 딜사이트. 필자 계산 순차입금(15.9조)과의 차이는 산정 범위 차이로 추정", section="그룹 재무 및 시사점")

# Checklist & messages
s = new_slide()
header(s, "그룹 재무 및 시사점", "다음 분기 체크포인트")
headline(s, "숫자로 약속하지 않은 목표는 검증할 수 없다. 투자자가 다음 분기에 확인할 다섯 가지와 CEO 메시지")
col_title(s, 0.3, 4.95, 1.4, "투자자 체크포인트")
col_title(s, 5.55, 4.98, 1.4, "부문별 CEO 메시지")
checks = [("재고효과를 뺀 정유 이익", "유가가 꺾이는 분기에 진짜 마진이 드러난다"),
          ("보상금·세액공제를 뺀 배터리 이익", "'26 2Q 흑자가 반복 가능한 이익인지 가르는 기준"),
          ("SK온 안에서 윤활유와 배터리 분리", "합쳐 보면 배터리 회복이 실제보다 좋아 보인다"),
          ("순차입금/EBITDA", "무디스의 '26년 2.7배 전망을 실제 숫자가 따라오는지"),
          ("추가 손상과 매각 일정", "SK지오센트릭 해외 법인, SK온 해외 공장")]
for i, (t_, d_) in enumerate(checks):
    y = 1.98 + i * 0.95
    rect(s, 0.3, y, 0.42, 0.42, fill=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb(s, 0.3, y, 0.42, 0.42, [[(str(i + 1), F_X, 14, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
    tb(s, 0.85, y - 0.03, 4.4, 0.8, [{"runs": [(t_, F_B, 11, INK, True)], "space": 2}, [(d_, F_R, 9.5, GRAY, False)]])
msgs = [("석유", "유가가 준 이익은 유가가 가져간다. 우리가 통제하는 마진만 우리 실적이다"),
        ("화학", "버티는 회사가 아니라, 줄여서 사는 회사가 돼야 한다"),
        ("윤활유", "좋은 사업의 현금은 공짜가 아니다. 쓸 곳과 한도를 먼저 정한다"),
        ("배터리", "보조금 없이 남는 1원이 진짜 흑자다"),
        ("E&S", "방파제를 성장엔진이라 부르려면, 숫자로 된 이정표가 필요하다"),
        ("그룹", "리밸런싱의 끝은 매각 공시가 아니라, 주주 몫 이익이 늘어난 손익계산서다")]
t = s.shapes.add_table(len(msgs), 2, Inches(5.55), Inches(1.98), Inches(4.98), Inches(0.72 * len(msgs))).table
t.columns[0].width = Inches(0.95); t.columns[1].width = Inches(4.03)
for i, (a, b) in enumerate(msgs):
    t.rows[i].height = Inches(0.72)
    for j, (v, face, col, fillc) in enumerate([(a, F_X, WHITE, ACCENT), (b, F_B, INK, WHITE if i % 2 else FILL)]):
        c = t.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = fillc
        c.vertical_anchor = MSO_ANCHOR.MIDDLE; c.margin_left = c.margin_right = Inches(0.08)
        p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
        p.alignment = PP_ALIGN.CENTER if j == 0 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = v; set_font(r, face, 10.5 if j == 0 else 10, col, True)
source(s, "공통 축: 회사가 목표로만 말하고 숫자로 공시하지 않는 항목(정제마진, 보조금 제외 이익, 발전 용량, 순차입금 목표, 부문별 ROIC)을 측정 가능한 약속으로 바꿀 것")
pageno(s)

# Sources
s = new_slide()
header(s, "참고", "출처 및 유의사항")
srcs = [
 "SK이노베이션·SK에너지·SK온·SK지오센트릭 2025 사업보고서(DART, 2026.3.16 제출)",
 "Investing.com · SK이노베이션 2026년 2분기 실적 컨퍼런스콜 / 에너지데일리 · 2분기 실적 / 더트래커 · 2026 상반기 정유·석화 리포트",
 "머니투데이 · 장용호·추형욱 신년사(2026.1.2) / 글로벌이코노믹 · 제19기 정기주주총회(2026.3.24) / 헤럴드경제 · ‘전기화 포지셔닝 원년’",
 "이투데이·데일리안 · SK온 2026년 ESS 20GWh 수주 목표 / 머니투데이 · ESS 목표 달성 ‘긍정적’(2026.10) / 한국금융신문 · 미국 ESS 1.5조 수주",
 "머니투데이·이투데이·문화일보 · 포드 합작 종결, SK온 테네시 출범(2026.5)",
 "이투데이 · 무디스 전망 ‘안정적’ 상향 / 블로터 · 신용등급 괴리 / 한국신용평가(2025.11)",
 "파이낸스스코프 · 공개 컨퍼런스콜(체질 정상화·무배당) / 딜사이트 · 순차입금 감축 / 뉴스웨이 · 자산 매각 / 마이데일리 · 울산 석화 재편 지연",
 "디일렉 · SK온 상장 의무 해소와 2차 리밸런싱 / 블로터 · SK온, SK엔무브 흡수",
]
tb(s, 0.3, 1.0, 10.2, 4.4, [{"runs": [(x, F_R, 10, INK, False)], "bullet": "•", "space": 6} for x in srcs])
rect(s, 0.3, 5.55, 10.23, 1.25, fill=FILL)
tb(s, 0.45, 5.6, 9.95, 1.15, [
    {"runs": [("유의사항", F_B, 10, INK, True)], "space": 3},
    {"runs": [("본 자료는 공시와 출처가 명시된 공개 자료를 바탕으로 한 분석이며 특정 증권의 매수·매도 권유가 아닙니다. 2026년 수치는 회사 발표를 인용한 보도와 증권사 자료로 사업보고서로 확정되지 않았습니다. 목표 판정과 CEO 과제는 작성자의 견해입니다.", F_R, 9, GRAY, False)]},
], anchor=MSO_ANCHOR.MIDDLE)
pageno(s)

prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
